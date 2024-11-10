from flask import Flask, request, jsonify
import spacy
import re
import pdfplumber
import docx
import openpyxl
import os
import mysql.connector
from mysql.connector import Error

# Cargar el modelo de clasificación y el modelo de reconocimiento de patrones
modelo_clasificacion = spacy.load("modelo_clasificacion")  # Asegúrate de haber entrenado y guardado tu modelo de clasificación
modelo_patrones = spacy.load("es_core_news_sm")

app = Flask(__name__)

# Configuración de la base de datos
db_config = {
    'host': 'localhost',
    'user': 'admin',        # Cambia a tu usuario de MySQL
    'password': '1234',  # Cambia a tu contraseña de MySQL
    'database': 'ClasificadorDocumentos'
}

# Función para conectar a la base de datos
def conectar_db():
    try:
        connection = mysql.connector.connect(**db_config)
        return connection
    except Error as e:
        print("Error al conectar con MySQL:", e)
        return None

# Función para clasificar el tipo de documento
def clasificar_documento(texto):
    doc = modelo_clasificacion(texto)
    scores = doc.cats
    tipo_documento = max(scores, key=scores.get)
    return tipo_documento

# Función para extraer patrones clave (fechas y nombres)
def extraer_patrones(texto):
    doc = modelo_patrones(texto)
    patrones = {
        "fechas": [],
        "nombres": []
    }
    # Regex para fechas
    regex_fecha = r'\b\d{1,2}/\d{1,2}/\d{2,4}\b|\b\d{1,2}-\d{1,2}-\d{2,4}\b'
    patrones["fechas"] = re.findall(regex_fecha, texto)
    # Nombres propios usando el modelo de spaCy
    for entidad in doc.ents:
        if entidad.label_ == "PER":  # "PER" es la etiqueta de persona en spaCy
            patrones["nombres"].append(entidad.text)

    return patrones

# Función para extraer texto de archivos PDF, DOCX, XLSX y TXT
def extraer_texto(archivo):
    contenido = ""
    nombre, extension = os.path.splitext(archivo.filename)

    if extension == '.pdf':
        with pdfplumber.open(archivo) as pdf:
            for pagina in pdf.pages:
                contenido += pagina.extract_text()

    elif extension == '.docx':
        doc = docx.Document(archivo)
        contenido = "\n".join([p.text for p in doc.paragraphs])

    elif extension == '.txt':
        contenido = archivo.read().decode('utf-8')

    elif extension == '.xlsx':
        wb = openpyxl.load_workbook(archivo)
        sheet = wb.active
        contenido = "\n".join(["\t".join([str(cell.value) for cell in row]) for row in sheet.iter_rows()])

    return contenido

# Ruta para procesar documentos cargados
@app.route('/procesar_documento', methods=['POST'])
def procesar_documento():
    if 'file' not in request.files:
        return jsonify({"error": "No se encontró ningún archivo"}), 400

    archivo = request.files['file']
    nombre_documento = archivo.filename
    contenido = extraer_texto(archivo)
    
    if not contenido:
        return jsonify({"error": "No se pudo extraer texto del archivo."}), 400

    # Clasificar el documento
    tipo_documento = clasificar_documento(contenido)

    # Extraer patrones clave
    patrones = extraer_patrones(contenido)

    # Conectar a la base de datos y guardar el documento
    connection = conectar_db()
    if connection is None:
        return jsonify({"error": "No se pudo conectar a la base de datos"}), 500

    try:
        cursor = connection.cursor()
        palabras_clave = ", ".join(patrones["nombres"] + patrones["fechas"])  # Combina nombres y fechas para almacenarlas como palabras clave

        # Insertar el documento en la base de datos
        cursor.execute(
            "INSERT INTO documentos (nombre_documento, tipo_documento, contenido_texto, palabras_clave) VALUES (%s, %s, %s, %s)",
            (nombre_documento, tipo_documento, contenido, palabras_clave)
        )
        connection.commit()
        return jsonify({
            "message": "Documento procesado y guardado exitosamente",
            "tipo_documento": tipo_documento,
            "patrones": patrones
        }), 201

    except Error as e:
        return jsonify({"error": f"Error al guardar el documento en la base de datos: {str(e)}"}), 500

    finally:
        if connection.is_connected():
            cursor.close()
            connection.close()

if __name__ == '__main__':
    app.run(debug=True)
